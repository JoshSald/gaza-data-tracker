"""Gaza's hospitals, and what WHO has said about each one by name.

- data/lists/hospitals.json: the 36 hospitals in the list OCHA publishes on HDX, compiled by the Ministry of Health
  and the Health Cluster (dated 9 November 2023). Updated with `python -m ingest hospitals`.
- data/lists/hospital-status.json: statements about a named hospital, each quoted word for word from a WHO
  document, newest last. Nothing here is inferred: a hospital WHO has not named keeps no status. Entries are added
  with `add_status` (see scripts/hospital_mentions.py for finding candidate sentences) and never edited.

Statuses follow WHO's wording:
  partially_functioning  WHO says the hospital is partially functional/operational or has resumed partial operations
  not_functioning        WHO says it is non-functional, non-operational or out of service
  destroyed              WHO says it has been destroyed
  services_reported      WHO names services running there but does not state the hospital's overall status
"""
from __future__ import annotations

import io
import json

from ..archive import fetch, keep
from ..paths import DATA
from ..text import pdf_text
from ..verify import quote_in_document

LIST_URL = ("https://docs.google.com/spreadsheets/d/e/2PACX-1vQ5TVXUBpR3mYG7jsTup0xhppUVyTd9GjyrN-F06KoSlPEe9gwTtYskBKsxMs3kSVnDZn84"
            "UyldS0-k/pub?output=xlsx")
LIST_PAGE = "https://data.humdata.org/dataset/state-of-palestine-health-0"
LIST_FILE = DATA / "lists" / "hospitals.json"
STATUS_FILE = DATA / "lists" / "hospital-status.json"
STATUSES = ("partially_functioning", "not_functioning", "destroyed", "services_reported")


def parse_list(content: bytes) -> tuple[list[dict], str | None]:
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    rows = list(wb["Hospitals"].iter_rows(values_only=True))
    header = list(rows[1])
    col = {name: header.index(name) for name in ("Hospital", "Arabic Name", "Classification", "Governorate", "Number of beds (All)")}
    hospitals = []
    for r in rows[2:]:
        if not r[col["Hospital"]]:
            continue
        beds = r[col["Number of beds (All)"]]
        hospitals.append({
            "name": str(r[col["Hospital"]]).strip(),
            "ar_name": str(r[col["Arabic Name"]] or "").strip(),
            "type": str(r[col["Classification"]] or "").strip(),
            "governorate": str(r[col["Governorate"]] or "").strip(),
            "beds": int(beds) if isinstance(beds, (int, float)) else None,
        })
    note = {str(r[0]).strip(): r[1] for r in wb["Note"].iter_rows(values_only=True) if r and r[0]}
    return hospitals, _iso(note.get("Date"))


def _iso(value) -> str | None:
    """'Nov 9th 2023' (as the sheet writes it) or a date cell, as YYYY-MM-DD."""
    import calendar
    import datetime as dt
    import re

    if isinstance(value, (dt.date, dt.datetime)):
        return value.strftime("%Y-%m-%d")
    m = re.match(r"([A-Za-z]{3})[a-z]*\.? (\d{1,2})(?:st|nd|rd|th)? (\d{4})", str(value or "").strip())
    months = {m.lower()[:3]: i for i, m in enumerate(calendar.month_name) if m}
    if not m or m.group(1).lower() not in months:
        return None
    return dt.date(int(m.group(3)), months[m.group(1).lower()], int(m.group(2))).isoformat()


def run(dry_run: bool) -> list[str]:
    doc = fetch(LIST_URL)
    hospitals, dated = parse_list(doc.content)
    lines = [f"Hospital list: {len(hospitals)} hospitals (list dated {dated})."]
    if dry_run:
        return lines + ["(dry run: nothing written)"]
    keep(doc)
    meta = {
        "source_note": "Compiled by the Gaza Ministry of Health and the Health Cluster; published by OCHA on HDX.",
        "document": {**doc.meta("OCHA (Humanitarian Data Exchange)"), "page": LIST_PAGE},
        "list_date": dated,
        "note": "Beds are the list's 'Number of beds (All)' before the war. The list does not give each hospital's status.",
    }
    LIST_FILE.parent.mkdir(parents=True, exist_ok=True)
    LIST_FILE.write_text(json.dumps({**meta, "hospitals": hospitals}, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    return lines


def load_statuses() -> list[dict]:
    return json.loads(STATUS_FILE.read_text())["statements"] if STATUS_FILE.exists() else []


def add_status(entries: list[dict]) -> list[dict]:
    """Append statements after checking each quote against the archived document. Returns the ones added."""
    hospitals = {h["name"] for h in json.loads(LIST_FILE.read_text())["hospitals"]}
    existing = load_statuses()
    seen = {(s["hospital"], s["document"]["sha256"], s["quote"]) for s in existing}
    added = []
    for e in entries:
        if e["hospital"] not in hospitals:
            raise ValueError(f"unknown hospital {e['hospital']!r}")
        if e["status"] not in STATUSES:
            raise ValueError(f"unknown status {e['status']!r}")
        path = e["document"]["archive"]
        text = pdf_text(open(path, "rb").read()) if path.endswith(".pdf") else open(path, encoding="utf-8").read()
        if not quote_in_document(e["quote"], text):
            raise ValueError(f"quote not found in {path}: {e['quote'][:80]}")
        if (e["hospital"], e["document"]["sha256"], e["quote"]) in seen:
            continue
        added.append(e)
    statements = sorted(existing + added, key=lambda s: (s["as_of"], s["hospital"]))
    STATUS_FILE.write_text(json.dumps({
        "note": "Statements by WHO that name a hospital and its status, quoted word for word. The newest per hospital is shown on the site.",
        "statuses": {
            "partially_functioning": "WHO says the hospital is partially functional or operational, or has resumed partial operations.",
            "not_functioning": "WHO says the hospital is non-functional, non-operational or out of service.",
            "destroyed": "WHO says the hospital has been destroyed.",
            "services_reported": "WHO names services running at the hospital but does not state its overall status.",
        },
        "statements": statements,
    }, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    return added

